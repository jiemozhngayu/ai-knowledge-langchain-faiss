"""生成企业知识库助手 LangChain + FAISS 更新版 PRD。"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parent.parent
OUTPUT = PROJECT_DIR / "项目文档" / "企业知识库助手_LangChain_FAISS_PRD_更新版.docx"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 98, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CONTENT_WIDTH = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, size=11, bold=None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_text(doc, text: str, *, bold=False, color=None, after=6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, 11, bold, color)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(item)
        set_run_font(run, 11)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(item)
        set_run_font(run, 11)


def add_heading(doc, text: str, level=1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 13 if level == 2 else 12, True, BLUE if level < 3 else DARK_BLUE)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_GRAY)
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, 10.5, True, INK)
    mark_header_row(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(value)
            set_run_font(run, 10.2)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for level, size, before, after, color in [
        (1, 16, 16, 8, BLUE),
        (2, 13, 12, 6, BLUE),
        (3, 12, 8, 4, DARK_BLUE),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("企业知识库助手｜LangChain + FAISS 更新版 PRD")
    set_run_font(header_run, 8.5, False, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("学习版设计文档｜2026-09-01")
    set_run_font(footer_run, 8.5, False, MUTED)


def build() -> None:
    doc = Document()
    configure_document(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("企业知识库助手")
    set_run_font(run, 25, True, INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    subrun = subtitle.add_run("LangChain + FAISS 本地部署方案｜产品需求文档（更新版）")
    set_run_font(subrun, 14, False, MUTED)

    metadata = doc.add_table(rows=4, cols=2)
    metadata.style = "Table Grid"
    set_table_geometry(metadata, [2700, 6660])
    entries = [
        ("文档版本", "V2.0（在原学习版 PRD 基础上更新）"),
        ("更新时间", "2026-09-01"),
        ("适用范围", "约 100 人规模的中小企业本地知识库学习与原型验证"),
        ("当前状态", "管理员资料上传、FAISS 入库、SQLite 元数据、检索验证已实现；正式认证未完成"),
    ]
    for row, (label, value) in zip(metadata.rows, entries):
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        row.cells[0].text = ""
        row.cells[1].text = ""
        set_run_font(row.cells[0].paragraphs[0].add_run(label), 10.5, True, INK)
        set_run_font(row.cells[1].paragraphs[0].add_run(value), 10.5)
    mark_header_row(metadata.rows[0])

    add_heading(doc, "1. 产品概述")
    add_text(doc, "企业知识库助手是一个本地部署的内部问答原型。员工通过聊天查询制度、流程和产品资料；系统先依据访问范围检索企业资料，再将有权限的内容交给本机大模型生成回答，并展示资料来源。更新版在原有问答基础上补齐了管理员资料入库闭环。")
    add_text(doc, "本项目的目标是教学、原型验证和本地方案演示，而不是直接作为生产系统上线。所有原始资料、向量索引、SQLite 元数据与本机模型均保存在 D 盘项目目录中。", bold=True, color=DARK_BLUE)

    add_heading(doc, "2. 背景与问题")
    add_bullets(doc, [
        "企业制度、报销流程、产品资料分散在 PDF、Word、Excel 和文本文件中，员工查找成本高。",
        "重复咨询占用 HR、财务、IT 与业务人员时间，且口头答复难以追溯。",
        "不同部门资料的保密要求不同，不能让普通员工默认看到财务、管理层内容。",
        "手工放入文件夹后再运行脚本入库，不利于管理员持续维护知识库。",
    ])

    add_heading(doc, "3. 产品目标与边界")
    add_table(doc, ["类别", "定义"], [
        ["产品目标", "让员工以自然语言查询有权限的企业资料；让管理员可上传资料、指定范围、自动向量化并验证检索结果。"],
        ["本期范围", "本机 Ollama、LangChain、FAISS、SQLite、FastAPI、Streamlit；支持 PDF/DOCX/XLSX/TXT。"],
        ["不在本期范围", "企业 SSO、真实员工目录同步、外网访问、云端模型、ERP/CRM 写操作、自动审批和多租户。"],
        ["上线边界", "当前演示身份和 X-Demo-Role 仅用于学习；不能作为正式权限认证。"],
    ], [1900, 7460])

    add_heading(doc, "4. 用户角色与权限")
    add_table(doc, ["角色", "当前学习版能力", "正式版要求"], [
        ["普通员工", "查询全员资料与本人部门资料；查看答案来源。", "从登录令牌读取身份与部门，服务端校验。"],
        ["部门员工", "如财务员工，可查询全员与财务部资料。", "细化到部门、岗位、资料密级与数据行权限。"],
        ["系统管理员", "可查看资料目录、上传资料、重建索引、查看日志、测试 FAISS 检索。", "真实管理员角色、MFA、操作审计、最小权限和审批。"],
    ], [1600, 3900, 3860])

    add_heading(doc, "5. 核心业务流程")
    add_heading(doc, "5.1 员工问答流程", 2)
    add_numbered(doc, [
        "员工在网页输入问题，页面提交问题与当前学习演示身份。",
        "后端根据角色、部门和资料访问范围过滤候选片段。",
        "FAISS 按向量距离检索；无关资料或“未授权资料明显更接近”的情况会直接拒答。",
        "只有通过权限过滤的片段会进入本机 Ollama 问答模型。",
        "页面展示回答和文件/页码来源；无可访问资料时显示统一提示。",
    ])
    add_heading(doc, "5.2 管理员资料入库流程", 2)
    add_numbered(doc, [
        "管理员选择资料访问范围：全员、财务部或管理层。",
        "上传 PDF、DOCX、XLSX、TXT，单文件最大 20 MB；同名文件直接拒绝，避免静默覆盖。",
        "后端将资料保存到 documents 对应范围目录，并写入 SQLite 文件目录。",
        "LangChain 读取资料，按 chunk_size=600、chunk_overlap=100 切分，再调用 qwen3-embedding:0.6b 生成向量。",
        "FAISS 先在临时目录创建完整新索引，成功后才切换为当前索引；失败时保留旧索引。",
        "管理员可在后台查看资料状态、片段数、操作日志，并使用“仅检索”验证结果。",
    ])

    add_heading(doc, "6. 功能需求")
    add_table(doc, ["编号", "功能", "需求说明", "验收要点"], [
        ["FR-01", "智能问答", "基于本机知识库回答，并展示文件来源与页码。", "公开制度问题能返回答案和来源。"],
        ["FR-02", "检索前权限", "检索结果先按全员/部门/管理员范围过滤，再给模型。", "普通员工提问财务付款审批不返回财务内容。"],
        ["FR-03", "管理员上传", "支持多文件上传、范围选择、格式与大小校验、重复文件拒绝。", "重复上传返回冲突提示，原文件不被覆盖。"],
        ["FR-04", "自动向量化", "上传后自动切分并构建 FAISS 索引。", "上传成功后片段数与资料目录同步增加。"],
        ["FR-05", "资料目录", "SQLite 保存来源、范围、大小、摘要状态、片段数和更新时间。", "后台可读取资料清单。"],
        ["FR-06", "索引维护", "管理员可对当前 documents 目录手工重建索引。", "新索引成功才切换；失败可保留旧索引。"],
        ["FR-07", "检索验证", "管理员可直接查看 FAISS 来源、范围、距离和片段预览。", "不调用问答模型即可定位资料是否入库。"],
        ["FR-08", "操作日志", "记录上传、索引重建与失败信息。", "后台展示最近操作记录。"],
    ], [850, 1500, 4200, 2810])

    add_heading(doc, "7. 页面与接口设计")
    add_heading(doc, "7.1 页面")
    add_table(doc, ["页面", "主要内容"], [
        ["问答首页", "身份演示选择、服务健康状态、对话输入、回答、来源展开。"],
        ["管理员后台", "资料数量/片段数/索引状态、上传区、手工重建、FAISS 检索验证、SQLite 资料目录、操作日志。"],
    ], [2500, 6860])
    add_heading(doc, "7.2 已实现接口")
    add_table(doc, ["接口", "用途", "权限说明"], [
        ["GET /api/health", "查看 Ollama 与 FAISS 索引状态。", "学习版公开健康检查。"],
        ["POST /api/chat", "按身份范围检索并生成回答。", "当前身份由请求传入，仅作演示。"],
        ["GET /api/admin/summary", "查看资料数量、片段数、索引状态、日志。", "学习版要求 X-Demo-Role: admin。"],
        ["GET /api/admin/documents", "读取 SQLite 资料目录。", "同上。"],
        ["POST /api/admin/documents/upload", "上传资料并自动重建索引。", "同上。"],
        ["POST /api/admin/reindex", "按现有资料重建索引。", "同上。"],
        ["GET /api/admin/search", "管理员验证 FAISS 检索，不调用模型。", "同上。"],
    ], [3250, 3600, 2510])

    add_heading(doc, "8. 技术架构与数据设计")
    add_text(doc, "技术栈：Streamlit 负责学习版网页；FastAPI 提供接口；LangChain 负责文档对象、切分与嵌入客户端；FAISS 保存本地向量和片段元数据；SQLite 保存文件目录与日志；Ollama 运行 qwen3-embedding:0.6b（向量）和 qwen2.5:3b（问答）。")
    add_table(doc, ["存储", "内容", "定位"], [
        ["documents/", "原始 PDF、DOCX、XLSX、TXT，按访问范围分目录。", "原始资料仓库，不被索引重建覆盖。"],
        ["data/faiss_index/", "当前可用的向量索引。", "语义相似度检索，不是关系型数据库。"],
        ["data/faiss_index_previous/", "上一次索引切换前保留的索引。", "构建失败时的恢复保护。"],
        ["data/knowledge.db", "资料元数据与操作日志。", "SQLite 目录库，不承担向量相似度计算。"],
    ], [2400, 4200, 2760])

    add_heading(doc, "9. 安全、权限与风险")
    add_bullets(doc, [
        "当前 X-Demo-Role 请求头可以被手工伪造，因此只能用于本机教学。正式版必须接入真实登录、JWT 或会话校验，并由服务端取得部门和角色。",
        "管理员上传需要增加病毒扫描、文件类型内容校验、敏感信息识别、容量配额、版本号和审批流程。",
        "FAISS 的 index.pkl 只允许加载本系统自行创建的文件，不应加载来源不明的索引文件。",
        "向量距离阈值 1.20 和未授权距离差 0.20 仅基于当前学习资料校准；生产前必须用企业真实问答集验证召回率、拒答率和权限越界率。",
        "本机模型和资料不上传云端，但同一局域网部署、操作系统账户权限和本地备份仍需要单独加固。",
    ])

    add_heading(doc, "10. 非功能需求")
    add_table(doc, ["维度", "学习版要求", "正式版演进方向"], [
        ["性能", "小型资料集下完成检索与本机问答；索引重建期间旧索引仍可使用。", "建立基准测试、并发队列、缓存和容量监控。"],
        ["可靠性", "新索引先构建后切换；失败时保留旧索引。", "定期备份、恢复演练、版本化与健康告警。"],
        ["可维护性", "上传、索引、目录、检索验证均可在管理员页操作。", "配置中心、日志检索、自动化测试和部署脚本。"],
        ["可审计性", "记录上传和重建动作。", "记录用户、IP、请求、资料版本、授权决定与导出行为。"],
    ], [1600, 3880, 3880])

    add_heading(doc, "11. 当前验收结果")
    add_table(doc, ["验证项", "结果"], [
        ["管理员上传", "已上传管理员示例 TXT，资料总数变为 5，知识片段总数变为 10。"],
        ["重复保护", "重复上传同名资料返回 409，未覆盖已存在文件。"],
        ["管理员接口", "管理员摘要、资料目录、检索接口均可实际调用。"],
        ["接口权限提示", "未携带管理员标识访问管理接口返回 403；但这只是学习版限制，不等于安全认证。"],
        ["权限问答", "普通员工查询超过 50000 元付款审批时，返回无当前可访问规定，不返回财务资料。"],
        ["新资料检索", "管理员可检索到新上传示例资料及其来源。"],
        ["页面验证", "Streamlit 管理员后台可显示资料数量、片段数和索引状态，无页面运行异常。"],
    ], [2500, 6860])

    add_heading(doc, "12. 后续迭代建议")
    add_numbered(doc, [
        "接入真实登录和 JWT：取消网页传入部门/角色的学习演示模式。",
        "实现资料版本与回收站：管理员删除资料时可恢复，避免直接删除原始文件。",
        "增加文件解析预览、失败原因与重复内容检测，减少上传后不可检索的问题。",
        "建立企业问答测试集，按部门分别评估检索命中、拒答、权限与回答质量。",
        "增加定时备份、恢复演练和局域网部署权限控制；再考虑 ERP、CRM、OA、飞书等连接。",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
