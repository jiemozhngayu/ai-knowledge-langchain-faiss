from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parent / "企业知识库助手_LangChain_FAISS_PRD.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_run_font(run, size=None, bold=None, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.1


def set_table_borders(table, color="D0D5DD"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_run_font(run, size=9, color=MUTED)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("企业知识库助手 · 产品需求文档")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("内部学习资料  |  第 ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(p)
    run = p.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("产品需求文档（PRD）")
    set_run_font(run, size=23, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("企业知识库助手 · LangChain + FAISS 本地方案")
    set_run_font(run, size=14, color=DARK_BLUE)

    metadata = doc.add_table(rows=4, cols=2)
    set_table_geometry(metadata, [1800, 7560])
    set_table_borders(metadata)
    rows = [
        ("文档版本", "V1.0"),
        ("编制日期", "2026年9月1日"),
        ("产品定位", "面向约100人中小企业的本地知识库智能问答学习版"),
        ("当前状态", "核心检索、权限过滤、后端接口和网页问答已验证"),
    ]
    for row, (label, value) in zip(metadata.rows, rows):
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(label_run, size=10.5, bold=True, color=INK)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, size=10.5, color="344054")

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [9360])
    set_table_borders(callout, color="B7CBE4")
    set_cell_shading(callout.cell(0, 0), LIGHT_BLUE)
    p = callout.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("核心结论：本方案将现有的 JSON + NumPy 学习检索升级为 LangChain + FAISS 本地向量检索，同时保留本地 Ollama、资料来源追溯与检索前权限过滤。")
    set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        size, color = 16, BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        size, color = 13, BLUE
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        size, color = 12, DARK_BLUE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, size=11, bold=True, color=INK)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run, size=11, color="344054")
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color="344054")
    return p


def add_data_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_GRAY)
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=10, bold=True, color=INK)
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=10, color="344054")
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    add_title(doc)

    add_heading(doc, "1. 产品背景与目标")
    add_body(doc, "企业内部制度、报销流程、产品资料和业务规范通常分散在 PDF、Word、Excel 或文本文件中。员工依赖人工询问或逐份翻找资料，导致重复咨询、资料难以定位和知识传承不足。")
    add_body(doc, "本项目面向约100人规模的中小企业，建设一个本地部署的知识库问答学习版。员工通过聊天提出问题，系统在其允许访问的资料范围内检索内容，并输出简洁答案与来源文件、页码。")
    add_body(doc, "目标：在不调用云端大模型 API 的前提下，形成可学习、可演示、可逐步迁移到正式系统的本地 RAG 架构。")

    add_heading(doc, "2. 产品范围")
    add_heading(doc, "2.1 本期范围", level=2)
    add_body(doc, "资料入库：支持 PDF、DOCX、XLSX、TXT，资料按 documents 文件夹下的第一层目录标记访问范围，例如“全员”“财务部”“管理层”。")
    add_body(doc, "知识检索：使用 LangChain 文档对象和 RecursiveCharacterTextSplitter 完成切分，当前参数为 chunk_size=600、chunk_overlap=100；使用 qwen3-embedding:0.6b 生成 1024 维本地向量，并使用 FAISS 保存索引。")
    add_body(doc, "问答与追溯：检索结果先完成相关度和权限过滤，再作为上下文交给本机问答模型；页面展示答案和来源文件、页码。")
    add_body(doc, "服务与页面：FastAPI 提供健康检查和问答接口，Streamlit 提供本机学习页面，默认使用 8001 与 8503 端口。")
    add_heading(doc, "2.2 不在本期范围", level=2)
    add_body(doc, "不包含 ERP、CRM、OA 写入操作，不执行自动审批，不支持对外公网发布，也不将网页演示身份视为正式登录与权限控制。")

    add_heading(doc, "3. 用户与权限设计")
    add_data_table(doc, ["角色/身份", "可访问范围", "本期说明"], [
        ("普通员工", "全员资料", "可查询公开制度和流程；财务内部资料不会进入模型上下文。"),
        ("财务员工", "全员资料 + 财务部资料", "可查询预算、付款审批等财务部授权资料。"),
        ("系统管理员", "全部资料", "学习版用于验证管理能力；正式版应从服务端 JWT 登录态获得身份。"),
    ], [1900, 2700, 4760])
    add_body(doc, "安全原则：正式系统不得信任网页提交的部门或角色字段。当前 Streamlit 身份选择仅为本机学习演示，后续必须接入服务端验证的 JWT，并从令牌中读取用户部门和角色。")

    add_heading(doc, "4. 核心业务流程")
    add_body(doc, "资料管理员将文件放入 documents 对应部门目录，执行入库脚本后，系统读取文件、保留文件名和页码、按600字符切分，并使用本地嵌入模型建立 FAISS 索引。")
    add_body(doc, "员工提问后，后端根据身份检索候选知识片段，先过滤超出相关度阈值或不允许访问的内容；仅将剩余资料发送给本机问答模型，最后返回答案和来源。")
    add_body(doc, "若没有找到当前账号可访问的相关资料，系统直接返回提示，不向问答模型传递未授权资料。")

    add_heading(doc, "5. 功能需求")
    add_heading(doc, "5.1 知识入库", level=2)
    add_body(doc, "系统应识别 PDF、DOCX、XLSX、TXT 文件；PDF 保留页码，Excel 以工作表为单位保留来源；资料元数据至少包含 source、page、access_scope。")
    add_body(doc, "当已有索引时，重建操作必须显式传入 --rebuild 参数，避免误覆盖。索引仅保存在新项目 data/faiss_index 目录，旧项目 embeddings.json 保持不变。")
    add_heading(doc, "5.2 检索与回答", level=2)
    add_body(doc, "系统应取相似候选结果并按距离阈值过滤。当前学习资料初步验证阈值为 1.0，后续需使用真实企业问题集重新校准。")
    add_body(doc, "问答模型当前默认 qwen2.5:3b。原因是本机 qwen3:4b 的当前模板会稳定输出较长思考过程，不适合直接显示给员工；qwen3-embedding:0.6b 继续作为向量模型。")
    add_heading(doc, "5.3 接口与页面", level=2)
    add_body(doc, "GET /api/health 返回 FAISS 索引和 Ollama 服务状态；POST /api/chat 接收问题与学习演示身份，并返回 answer、sources。网页应显示服务状态、聊天记录和可展开的资料来源。")

    add_heading(doc, "6. 技术架构与数据设计")
    add_data_table(doc, ["层级", "组件", "职责"], [
        ("资料层", "documents / data", "保存原始资料、FAISS 索引及后续日志与备份。"),
        ("知识处理层", "LangChain + TextSplitter", "统一文档对象、切分、元数据管理。"),
        ("检索层", "FAISS + 权限过滤", "本地相似度检索；先过滤权限，再向模型提供上下文。"),
        ("模型层", "Ollama", "本地嵌入模型和本地问答模型。"),
        ("服务与界面", "FastAPI + Streamlit", "提供接口、健康检查和学习演示网页。"),
    ], [1800, 2800, 4760])
    add_body(doc, "本机代理兼容：Python 调用嵌入模型时使用 127.0.0.1:11434 且设置 trust_env=False，避免系统代理错误转发本机请求并造成 502。")

    add_heading(doc, "7. 验收标准")
    add_data_table(doc, ["场景", "预期结果", "当前验证"], [
        ("索引建立", "4份学习资料可生成 FAISS 索引并可从磁盘加载。", "已完成：生成9个知识片段。"),
        ("公开制度检索", "普通员工提问住宿报销，返回500元/晚及来源。", "已完成。"),
        ("财务权限拦截", "普通员工提问超50000元付款审批，不显示财务资料。", "已完成。"),
        ("财务部门检索", "财务员工提问同一问题，返回财务负责人复核并提交总经理审批。", "已完成。"),
        ("服务与网页", "健康接口正常，网页可显示答案与资料来源。", "已完成。"),
    ], [2100, 4300, 2960])

    add_heading(doc, "8. 后续迭代计划")
    add_body(doc, "第一优先级：将旧项目 JWT 登录迁移到新服务，替换当前演示身份字段；同时将用户、部门、启停状态和审计日志纳入新项目数据层。")
    add_body(doc, "第二优先级：新增管理员资料上传、重建索引、索引版本、用户搜索与文档管理页面，并为错误导入、重复文件和索引重建增加可恢复日志。")
    add_body(doc, "第三优先级：建立真实问题集，评估召回率、答案依据、响应速度和权限绕过风险；根据结果校准 chunk_size、chunk_overlap、检索数量和距离阈值。")

    add_heading(doc, "9. 风险与边界")
    add_body(doc, "本机部署需要保持 Ollama 服务与模型可用；模型回答只能作为资料辅助查询，不能替代制度原文、管理审批或法律财务决策。")
    add_body(doc, "FAISS 本地索引会生成 index.pkl，只能加载本项目自己生成、可信来源的索引文件；不得加载来源不明的索引。")
    add_body(doc, "资料权限的关键控制点是“进入问答模型前过滤”。即使检索层返回候选内容，未授权片段也不得拼接到模型提示词或页面来源中。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
