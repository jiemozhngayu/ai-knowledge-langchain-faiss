"""读取资料、切分文本并原子化更新本地 FAISS 索引。"""

from __future__ import annotations

import shutil
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from docx import Document as WordDocument
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pypdf import PdfReader

from faiss_config import DOCUMENTS_DIR, INDEX_DIR, get_embeddings


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".txt"}
CHUNK_SIZE = 600
CHUNK_OVERLAP = 100
BUILDING_INDEX_DIR = INDEX_DIR.with_name("faiss_index_building")
PREVIOUS_INDEX_DIR = INDEX_DIR.with_name("faiss_index_previous")


@dataclass(frozen=True)
class IndexBuildResult:
    chunk_count: int
    file_chunk_counts: dict[str, int]


def get_access_scope(file_path: Path) -> str:
    """根据 documents 下第一层目录确定访问范围。"""
    relative = file_path.relative_to(DOCUMENTS_DIR)
    return relative.parts[0] if len(relative.parts) > 1 else "全员"


def read_file(file_path: Path) -> list[Document]:
    """读取 PDF、DOCX、XLSX 或 TXT，并保留检索所需元数据。"""
    scope = get_access_scope(file_path)
    source = str(file_path.relative_to(DOCUMENTS_DIR))
    suffix = file_path.suffix.lower()
    pages: list[tuple[int, str]] = []

    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        pages = [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, start=1)]
    elif suffix == ".docx":
        word = WordDocument(str(file_path))
        content = "\n".join(paragraph.text for paragraph in word.paragraphs if paragraph.text.strip())
        pages = [(1, content)]
    elif suffix == ".xlsx":
        workbook = load_workbook(file_path, data_only=True)
        for number, sheet in enumerate(workbook.worksheets, start=1):
            rows = [" | ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True)]
            pages.append((number, f"工作表：{sheet.title}\n" + "\n".join(row for row in rows if row.strip())))
    elif suffix == ".txt":
        pages = [(1, file_path.read_text(encoding="utf-8"))]

    display_source = source
    if file_path.name.startswith("飞书文档_") and suffix == ".txt":
        for _, text in pages:
            for line in text.splitlines()[:5]:
                if line.startswith("飞书文档标题："):
                    title = line.removeprefix("飞书文档标题：").strip()
                    if title:
                        display_source = f"飞书文档：{title}"
                    break

    return [
        Document(
            page_content=text.strip(),
            metadata={
                "source": source,
                "display_source": display_source,
                "page": page,
                "access_scope": scope,
            },
        )
        for page, text in pages
        if text.strip()
    ]


def collect_documents() -> list[Document]:
    """收集并切分全部受支持资料。"""
    raw_documents: list[Document] = []
    for file_path in DOCUMENTS_DIR.rglob("*"):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_SUFFIXES:
            raw_documents.extend(read_file(file_path))
    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    return splitter.split_documents(raw_documents)


def rebuild_index() -> IndexBuildResult:
    """先在临时目录建完索引，再替换当前索引，降低中途失败破坏旧索引的风险。"""
    chunks = collect_documents()
    if not chunks:
        raise RuntimeError("没有可入库的资料。请上传 PDF、DOCX、XLSX 或 TXT 文件。")

    if BUILDING_INDEX_DIR.exists():
        shutil.rmtree(BUILDING_INDEX_DIR)
    vector_store = FAISS.from_documents(chunks, get_embeddings())
    vector_store.save_local(str(BUILDING_INDEX_DIR))

    if PREVIOUS_INDEX_DIR.exists():
        shutil.rmtree(PREVIOUS_INDEX_DIR)
    if INDEX_DIR.exists():
        INDEX_DIR.replace(PREVIOUS_INDEX_DIR)
    try:
        BUILDING_INDEX_DIR.replace(INDEX_DIR)
    except Exception:
        if PREVIOUS_INDEX_DIR.exists() and not INDEX_DIR.exists():
            PREVIOUS_INDEX_DIR.replace(INDEX_DIR)
        raise

    counts = Counter(str(chunk.metadata["source"]) for chunk in chunks)
    return IndexBuildResult(chunk_count=len(chunks), file_chunk_counts=dict(counts))
